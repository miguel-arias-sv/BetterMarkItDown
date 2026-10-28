"""Word and PowerPoint files. Offline: the vision model and Office are faked."""

from __future__ import annotations

import io
from pathlib import Path

import pytest

from bettermarkitdown import core, formats
from bettermarkitdown.config import Options


def _chart_png(width: int = 300, height: int = 200) -> bytes:
    pymupdf = pytest.importorskip("pymupdf")
    doc = pymupdf.open()
    page = doc.new_page(width=width, height=height)
    page.draw_rect(pymupdf.Rect(20, 40, 80, height - 20), color=(0, 0, 0), fill=(0.3, 0.5, 0.8))
    png = page.get_pixmap(dpi=72).tobytes("png")
    doc.close()
    return png


@pytest.fixture
def report_docx(tmp_path: Path) -> Path:
    docx = pytest.importorskip("docx")
    document = docx.Document()
    document.add_heading("Harvest report", level=1)
    document.add_paragraph("Yields rose in every region.")
    document.add_picture(io.BytesIO(_chart_png()), width=docx.shared.Inches(3))
    document.add_picture(io.BytesIO(_chart_png(20, 20)))           # an icon: no vision call
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Region", "Yield"
    table.cell(1, 0).text, table.cell(1, 1).text = "North", "3.6"
    path = tmp_path / "My report.docx"
    document.save(path)
    return path


@pytest.fixture
def deck_pptx(tmp_path: Path) -> Path:
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches
    deck = pptx.Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Harvest yields"
    slide.shapes.add_picture(io.BytesIO(_chart_png()), Inches(1), Inches(2), width=Inches(4))
    path = tmp_path / "deck.pptx"
    deck.save(path)
    return path


@pytest.fixture
def fake_vision(monkeypatch):
    calls = []

    def describe(client, model, png, prompt, **kwargs):
        calls.append(prompt)
        return "**[FIGURE]** A bar chart.\n\n| Region | Yield |\n|---|---|\n| North | 3.6 |"

    monkeypatch.setattr(core, "describe_image", describe)
    return calls


def _options(source: Path, tmp_path: Path, **extra) -> Options:
    return Options(pdf=source, output=tmp_path / "out" / (source.stem + ".md"),
                   work_dir=tmp_path / "work", base_url="http://localhost:1/v1",
                   model="m", **extra)


def test_word_pictures_are_saved_and_described_where_they_stood(report_docx, tmp_path,
                                                                 fake_vision):
    pytest.importorskip("mammoth")
    opts = _options(report_docx, tmp_path)
    result = formats.convert_office(opts)
    text = opts.output.read_text(encoding="utf-8")

    assert text.startswith("# My report\n")
    assert "data:image" not in text, "no picture is left as an empty placeholder"
    assert "| Region | Yield |" in text and "Harvest report" in text      # MarkItDown's part
    assert fake_vision == [formats.FIGURE_PROMPT]                         # the icon was skipped
    assert "**Picture 1 - transcribed by m:**" in text and "> **[FIGURE]** A bar chart." in text
    assert "](My%20report_assets/My%20report_picture_01.png)" in text      # link survives spaces
    assert (tmp_path / "out" / "My report_assets" / "My report_picture_02.png").is_file()
    assert text.index("Yields rose") < text.index("Picture 1") < text.index("| North | 3.6 |")
    assert (result.unit, result.pages_converted, result.pages_selected) == ("picture", 1, 1)


def test_powerpoint_pictures_are_described_too(deck_pptx, tmp_path, fake_vision):
    pytest.importorskip("pptx")
    opts = _options(deck_pptx, tmp_path)
    formats.convert_office(opts)
    text = opts.output.read_text(encoding="utf-8")
    assert "Harvest yields" in text and "A bar chart." in text and "data:image" not in text


def test_text_only_keeps_the_pictures_but_calls_no_model(report_docx, tmp_path, fake_vision):
    pytest.importorskip("mammoth")
    opts = _options(report_docx, tmp_path, no_vision=True)
    result = formats.convert_office(opts)
    text = opts.output.read_text(encoding="utf-8")
    assert fake_vision == [] and result.calls == 0
    assert "My%20report_picture_01.png" in text and "transcribed by" not in text


def test_a_picture_that_cannot_be_read_is_named_not_dropped():
    pytest.importorskip("PIL")
    png, size = formats._to_png(b"not an image at all")
    assert png is None and size == (0, 0)


def test_a_spending_cap_stops_cleanly_and_says_which_pictures_are_left(
        report_docx, tmp_path, monkeypatch):
    pytest.importorskip("mammoth")

    def broke(*_a, **_k):
        raise core.BudgetExceeded("cap")

    monkeypatch.setattr(core, "describe_image", broke)
    opts = _options(report_docx, tmp_path)
    result = formats.convert_office(opts)
    assert result.stopped_reason == "budget"
    assert "run it again to finish" in opts.output.read_text(encoding="utf-8")
    assert result.pages_selected - result.pages_converted == 1


# --- through PDF -------------------------------------------------------------------

def test_office_is_tried_first_then_libreoffice(tmp_path, monkeypatch):
    source = tmp_path / "old.doc"
    source.write_bytes(b"\xd0\xcf\x11\xe0")
    dest = tmp_path / "work" / "old.pdf"
    ran = []

    def fake_run(command, env=None, timeout=300.0):
        ran.append(command[0])
        if command[0] == "powershell":
            Path(env["BMID_DEST"]).write_bytes(b"%PDF-1.4")
            return True
        return False

    monkeypatch.setattr(formats, "_run", fake_run)
    monkeypatch.setattr(formats, "office_installed", lambda app: app == "Word")
    assert formats.office_to_pdf(source, dest) == "Microsoft Word" and dest.is_file()

    dest.unlink()
    ran.clear()
    monkeypatch.setattr(formats, "office_installed", lambda app: False)
    monkeypatch.setattr(formats, "find_soffice", lambda: "soffice")

    def libreoffice(command, env=None, timeout=300.0):
        ran.append(command[0])
        outdir = Path(command[command.index("--outdir") + 1])
        (outdir / "old.pdf").write_bytes(b"%PDF-1.4")
        return True

    monkeypatch.setattr(formats, "_run", libreoffice)
    assert formats.office_to_pdf(source, dest) == "LibreOffice" and dest.is_file()
    assert ran == ["soffice"]


def test_with_no_converter_the_message_says_how_to_get_one(tmp_path, monkeypatch):
    source = tmp_path / "slides.ppt"
    source.write_bytes(b"\xd0\xcf\x11\xe0")
    monkeypatch.setattr(formats, "office_installed", lambda app: False)
    monkeypatch.setattr(formats, "find_soffice", lambda: "")
    with pytest.raises(formats.ConverterMissing, match="LibreOffice"):
        formats.office_to_pdf(source, tmp_path / "slides.pdf")


def _fake_office(monkeypatch):
    """Stand in for Word / PowerPoint: tests must never start the real programs."""
    made = []

    def fake_convert_file(src, dest, target="pdf"):
        made.append(dest.suffix)
        dest.parent.mkdir(parents=True, exist_ok=True)
        dest.write_bytes(b"%PDF-1.4" if target == "pdf" else b"PK")
        return "Microsoft PowerPoint"

    monkeypatch.setattr(formats, "office_convert", fake_convert_file)
    monkeypatch.setattr(formats, "office_to_pdf", lambda s, d: fake_convert_file(s, d, "pdf"))
    return made


def test_cli_saves_a_legacy_file_in_the_new_format_and_reads_it(tmp_path, monkeypatch, capsys):
    """Through PDF, tables come out as loose lines; through PPTX/DOCX they survive."""
    from bettermarkitdown import cli

    source = tmp_path / "lecture.ppt"
    source.write_bytes(b"\xd0\xcf\x11\xe0")
    made = _fake_office(monkeypatch)
    captured = {}

    def fake_office(opts, on_event=None, original=None):
        captured.update(opts=opts, original=original)
        raise KeyboardInterrupt

    monkeypatch.setattr(formats, "convert_office", fake_office)
    monkeypatch.chdir(tmp_path)
    assert cli.main([str(source), "--no-vision"]) == 130
    assert made == [".pptx"]
    opts = captured["opts"]
    assert opts.pdf.name == "lecture.pptx" and opts.pdf.parent.name == "lecture_work"
    assert captured["original"] == source
    assert opts.output == tmp_path / "lecture.md" and opts.title == "lecture"
    assert "with Microsoft PowerPoint" in capsys.readouterr().out


def test_as_pages_sends_a_file_through_pdf(tmp_path, monkeypatch):
    from bettermarkitdown import cli, core

    source = tmp_path / "lecture.ppt"
    source.write_bytes(b"\xd0\xcf\x11\xe0")
    made = _fake_office(monkeypatch)
    captured = {}

    def fake_convert(opts, on_event=None):
        captured["opts"] = opts
        raise KeyboardInterrupt

    monkeypatch.setattr(core, "convert", fake_convert)
    monkeypatch.chdir(tmp_path)
    assert cli.main([str(source), "--no-vision", "--as-pages"]) == 130
    assert made == [".pdf"] and captured["opts"].pdf.name == "lecture.pdf"


def test_cli_refuses_formats_it_cannot_read(tmp_path, capsys):
    from bettermarkitdown import cli

    other = tmp_path / "data.xyz"
    other.write_text("x")
    assert cli.main([str(other)]) == 1
    assert ".docx" in capsys.readouterr().err
